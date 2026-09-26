"""Append the v1326/v1327 delivery-parser release record to maintained DOCX files."""

from pathlib import Path

from docx import Document


ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs" / "maintenance"


def append_record(name: str, title: str, paragraphs: list[str]) -> None:
    target = DOCS / name
    document = Document(target)
    document.add_paragraph()
    heading = document.add_paragraph()
    heading.add_run(title).bold = True
    for text in paragraphs:
        document.add_paragraph(text)
    document.save(target)


append_record(
    "AI开发项目_项目说明文档.docx",
    "2026年9月27日 v1326 v1327 外卖嵌套商品名解析修复",
    [
        "两边共有：网页 v1326、私人内置网页 v1327、私人 iOS 1.0.397 (397)，原生桥 41。真实外卖动作现在使用按括号深度扫描的统一解析器，商品名包含【夜宵专享】等内部中文方括号时不再提前结束动作。",
        "聊天正文分离、缺动作修复结果和通话动作消费均复用同一解析结果。已识别外卖动作开头但未被完整消费的异常标签失败关闭，不执行，也不显示为角色聊天气泡。门店搜索、购物车、付款确认和真实支付边界未修改。",
        "发布核验同时覆盖网页与私人内置入口的半角外层、全角外层、动作前自然开场白、通话路径和畸形标签。自动化不创建真实订单、不执行支付；Mac 编译、签名和真实 iPhone 仍需另行验收。",
    ],
)

append_record(
    "AI开发项目_Bug记录模板.docx",
    "2026年9月27日 真实外卖商品名内部方括号导致动作截断",
    [
        "现象：角色生成肯德基四件套动作后，自动化可能停在主页面，后台得到的商品名不完整，或者控制标签残片出现在角色气泡。示例商品名为【夜宵专享】吃堡堡4件套。",
        "根因：旧实现用排除 ] 和 】 的正则寻找动作结尾，遇到商品名内部第一个 】 就提前结束，不能表达嵌套边界。503 是本机隧道连接状态问题，与本解析缺陷不是同一根因。",
        "解决：新增 deliveryStructuredActionTags、deliveryIsolateStructuredActions 和 deliveryReplaceStructuredActions，按外层括号种类与深度配对；聊天、修复和通话统一调用。检测到外卖标签开头但无法完整消费时静默拦截。",
        "验证：新增真实运行解析回归，覆盖网页和私人内置 app.js、半角与全角外层、嵌套商品名、聊天与通话消费、畸形标签失败关闭；并加入 permanent-fix-guard 与私人包反向标记。未运行真实下单或支付。",
    ],
)

append_record(
    "AI开发项目_Bug修改规范.docx",
    "2026年9月27日 结构化动作嵌套边界规则",
    [
        "结构化动作的字段值允许包含与业务名称有关的括号时，禁止使用“遇到第一个右括号就结束”的排除字符正则。必须按外层定界符种类和嵌套深度扫描，并让聊天、补判、重试和通话共用同一解析入口。",
        "已识别控制标签前缀但未完整解析或消费的内容必须失败关闭：不得执行残缺动作，不得把控制标签或残片显示为角色气泡。新增解析器必须同时覆盖网页和私人内置页，并以永久修复守卫防止后续整文件覆盖丢失。",
        "外卖可用性错误要拆分定位：HTTP 502/503/530 属于本地服务、隧道或上游连接层；门店或商品没有命中属于自动化页面层；动作字段被截断属于前端结构化协议层。不得用重启隧道掩盖解析缺陷，也不得用前端解析补丁冒充网络恢复。",
    ],
)
