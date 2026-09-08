"""Append release decisions without rewriting any historical maintenance record."""
from pathlib import Path
from docx import Document
ROOT=Path(__file__).resolve().parents[1]
TITLE='2026年9月8日 v1211 北京时间与私人更新发布'
ROWS={
 'AI开发项目_项目说明文档.docx':[
  '用户确认本次仅智能家居新改动不进入网页；两边共有已完成功能照常推送网页并进入私人包。私人后台AI拒绝说明隔离按此前授权仍仅私人。网页/内置页v1211，私人iOS1.0.336（336），桥36。未完成快捷指令后台草稿排除在提交、包和部署之外。',
  '北京时间气泡默认关闭，开启后以实际保存时间换算UTC+8，不使用模型编造时间，缺失旧时间显示未知，关闭恢复旧显示。私人新增门锁SVG对齐和只交付解锁事件给角色；保留Face ID及真实回读。私人既有性能和原生恢复差异保留，不进入网页。',
  '私人仍为Mac待编译源码候选，未签名、未真机安装；角色发起解锁问题未取得本轮真机解决证据。线上部署与文件哈希结果见父目录发布结果_v1211_2026-09-08.md，不能用提交或包生成代替上线完成。'],
 'AI开发项目_Bug记录模板.docx':[
  '发布版本机械替换首次误及335色相与0.335花束比例和对应测试。已恢复原业务数值并收紧版本字段规则，公开/私人旧smart-home.js与gift-effects.js无业务差异。旧标题断言、当前安装指南及门锁源副本身份不一致逐项修复，历史iOS327记录和旧包保留。',
  '新身份全仓1636/1636通过；两边核心聊天16组合及输入、让TA回、自动查看、存档恢复，原文多场景及私人后台拒绝/动作隔离/诊断持久化通过。北京时间浏览器设置为洛杉矶时区，双方气泡仍显示正确UTC+8，关闭回退正常。包由已提交树归档，显式跟踪被私人通配忽略的北京时间模块。',
  '部署前函数列表版本与旧记录不一致；已下载线上phone-role-push与HEAD源码统一换行后完全一致。仅部署此函数的新私人标记保护，不运行快捷指令迁移或部署其他服务；部署后须再下载比对。详细检查见v1211_发布核验.md，实际发布结果单独记录。'],
 'AI开发项目_Bug修改规范.docx':[
  '版本更新不可使用无上下文的纯数字替换：iOS编号可能与RGB色相、尺寸比例和业务测试数字碰撞。只更新明确版本字段与资源查询参数；同时核验私人源副本一致和历史安装文档不被错误提升。',
  '未完成草稿可以留在原位，但发布脚本必须列出精确排除文件，仅从已提交树生成产物，不能用宽泛忽略允许任意未跟踪文件混入。后台列表版本与记录不符时应下载实际源码核对，既不能盲覆盖，也不能仅以旧记录声明当前线上状态。']}
for name,rows in ROWS.items():
 p=ROOT/'docs/maintenance'/name;d=Document(p);old=[x.text for x in d.paragraphs];assert TITLE not in old
 d.add_page_break();d.add_paragraph(TITLE,style='Heading 1')
 for text in rows:d.add_paragraph(text)
 d.save(p);assert [x.text for x in Document(p).paragraphs][:len(old)]==old
 print(name+': appended, history preserved')
