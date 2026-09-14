from pathlib import Path
from docx import Document

ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs/maintenance"
TITLE = "v1241 和 v1242 大存档备份流畅修复"
DETAIL = "详见 docs/maintenance/v1241_v1242_大存档备份流畅修复核验.md。"

updates = {
    "AI开发项目_项目说明文档.docx": [
        "公开网页升至 v1241，私人内置源码升至 v1242。私人 iOS 保持 1.0.356（356），原生桥保持 38，本轮不制作私人安装包。",
        "完整备份直接分段读取当前存档。数组按固定长度逐项处理，对象进入时固定键列表，不再为十万级数组预先建立完整键表，也不再复制整份状态。每 64 个工作单元检查一次 8 毫秒时间预算并让出主线程。",
        "116,736 条聊天和 48 张媒体的 WebKit、Edge 页面导出、导入、保存和重启恢复通过；Windows 全仓 1,831/1,831 通过。真实 iPhone 和华为原浏览器仍待线上更新后复验。" + DETAIL,
    ],
    "AI开发项目_Bug记录模板.docx": [
        "现场诊断：v1239 在 backup-copy 阶段连续出现 7,640、12,919 和 36,005 毫秒事件循环停顿，随后页面被销毁；没有进入序列化、图片读取或下载阶段。全部脚本同为 v1239，排除混合缓存。",
        "根因：旧备份先递归复制整份状态，并对每个数组调用 Object.keys。十万级导入存档会同时保留原状态、完整数组键表和复制状态；让出主线程发生在完整键表创建之后。",
        "修复与证据：取消整份复制，数组按下标分段遍历，路径复用并按时间预算让出。旧实现会被新增的六位数数组回归阻断，新实现完整导出 116,736 项；WebKit 最大事件循环间隔 195 毫秒，Edge 为 280 毫秒。" + DETAIL,
    ],
    "AI开发项目_Bug修改规范.docx": [
        "大存档卡顿必须按诊断阶段定位。若最后阶段是 backup-copy，且没有 backup-serialize、backup-image-read 或 backup-download，不得把故障归因于下载按钮、Service Worker 或单张损坏图片。",
        "移动浏览器备份不得先为大型数组建立完整键表或复制整个状态树。数组应固定长度后按下标分段读取；分段算法必须用十万级数据验证事件循环间隔、导出、导入、保存和重启恢复。",
        "MacIntel 不能单独判定设备是 Mac。iPhone Safari 可使用桌面级 User-Agent，应结合 standalone、视口和像素比判断。pagehide 证明文档被销毁，不足以单独证明系统强杀。" + DETAIL,
    ],
}

for name, paragraphs in updates.items():
    path = DOCS / name
    doc = Document(path)
    if any(p.text.strip() == TITLE for p in doc.paragraphs):
        raise RuntimeError(f"{name}: release section already exists")
    doc.add_page_break()
    doc.add_heading(TITLE, level=1)
    for paragraph in paragraphs:
        doc.add_paragraph(paragraph)
    doc.save(path)
    print(f"updated {name}")
