"""Append the private backup outer-timeout finding without rewriting history."""

from pathlib import Path

from docx import Document
from lxml import etree


ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs" / "maintenance"
TITLE = "私人每日云备份提交等待修复候选 2026年9月21日"
DETAIL = "完整代码证据与原始设计边界见 v1240_私人每日分块云备份核验.md。当前仅有 Windows 源码与自动测试证据；尚未完成 Mac 编译、签名、真实 iPhone 54 MB 上传、跨日自动备份或云端恢复验收。打包与推送状态以实际发布结果为准。"

UPDATES = {
    "AI开发项目_项目说明文档.docx": (
        "私人 v1274 源码及现有 Mac 源码包实际包含 private-cloud-backup.js、四个分片动作、原生文件上传和 Xcode 文件夹同步归属，因此不能把本次失败写成已证实的漏包。新检查发现分片提交存在两层网页等待：WKWebView 原生桥允许 account.backup.file.commit 等待 660 秒，但外层 privatePhoneAccountCall 仍沿用 120 秒默认值。54 MB 备份在慢网下可能由外层先报超时，原生上传即使仍在进行也无法写入当天成功标记。候选修复让提交调用显式等待 720 秒，覆盖原生 660 秒上限；其他账号、恢复、镜像和数据保护规则不变。"
    ),
    "AI开发项目_Bug记录模板.docx": (
        "用户确认当前私人安装为 v1274，手动备份仍失败，云端日期停在 2026年9月2日。历史截图中的 account.backup.upload 属于旧整对象路径，不能直接代表 v1274 当前动作；当前真机失败提示与诊断尚未取得。代码和包内容核对证明 v1274 没有遗漏分片组件，但 v1240 记录所称的 660 秒只覆盖 SmallPhoneNative.request，private-cloud-backup.js 经 privatePhoneAccountCall 调用 commit 时未传 timeoutMs，仍会在 120 秒先拒绝。这是可确定的超时层不一致，也是能解释大备份慢网失败的结构性缺陷；尚不能声称它是该真机本次失败的唯一原因。候选修改为 commit 显式 720 秒，分片大小、原生 180/600 秒网络时限、旧云备份保护和 saved=true 成功条件均保留。专项及永久门禁共 240/240 通过，Mac 与真机未验证。"
    ),
    "AI开发项目_Bug修改规范.docx": (
        "长原生任务必须逐层核对超时，不能只放宽 URLSession 或 SmallPhoneNative.request 就宣称网页会等够。调用链中的业务包装器、桥请求、原生网络请求和资源总时限必须满足外层大于等于内层，并用测试记录实际传入的 timeoutMs。私人每日云备份是永久发布门禁：每个私人版本必须同时保留 private-cloud-backup.js、两个入口引用、begin/chunk/commit/abort、PrivateBackupFileStore、fromFile 上传、Xcode Target 资源归属和提交阶段等待关系；任一缺失都阻止测试与打包。源码包包含文件不等于 Mac 构建或真机已运行该文件。"
    ),
}


def body_nodes(doc: Document) -> list[bytes]:
    return [
        etree.tostring(node)
        for node in doc.element.body
        if not node.tag.endswith("}sectPr")
    ]


for filename, body in UPDATES.items():
    path = DOCS / filename
    doc = Document(path)
    assert TITLE not in [paragraph.text for paragraph in doc.paragraphs], filename
    before = body_nodes(doc)
    doc.add_page_break()
    doc.add_heading(TITLE, 1)
    doc.add_paragraph(body)
    doc.add_paragraph(DETAIL)
    doc.save(path)

    check = Document(path)
    assert body_nodes(check)[: len(before)] == before, filename
    assert TITLE in [paragraph.text for paragraph in check.paragraphs], filename

    text_path = path.with_suffix(".txt")
    original = text_path.read_text(encoding="utf-8")
    assert TITLE not in original, text_path.name
    text_path.write_text(original.rstrip() + "\n\n" + TITLE + "\n" + body + "\n" + DETAIL + "\n", encoding="utf-8")
    print(filename + ": appended; previous body preserved")
