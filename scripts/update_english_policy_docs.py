"""Append the authorized policy change while preserving every historical paragraph."""
from pathlib import Path
from docx import Document

ROOT = Path(__file__).resolve().parents[1]
TITLE = '2026年9月8日 纯英文拦截与心情格式 未发布'
ROWS = {
    'AI开发项目_项目说明文档.docx': [
        '用户最终确认：纯英文无中文翻译拦截与心情标签格式要求在原文开关开启、关闭时都生效。网页和私人版的线上、约会、共同生活及配角、电话、后台生成和收件均覆盖。除此之外关闭恢复旧处理，开启保留原有原文行为。心情标签仅加强提示词，不更改解析器和诊断判断。空回复暂不修复。',
        '新后台源码尚未部署；本轮不提交、推送、打包或更改版本号。私人智能家居与快捷指令草稿不动。没有真实模型、APNs 或 iPhone 验收。完整范围及验证见 2026-09-08_纯英文拦截与心情格式_未发布.md。',
    ],
    'AI开发项目_Bug记录模板.docx': [
        '旧代码开启原文时能把纯英文角色正文送入聊天。双入口单元回归先失败；隔离浏览器旧基线 9defbd8b 实际显示英文，新代码同断言通过。英文门禁在角色请求返回后、正文和动作处理前生效；中文元数据不能充当译文，正常中文及带翻译外语保留。',
        '回归捕获私人约会绕过公共线下请求包装器，已对其直接 chatAPI 主请求和纠正请求补齐标记。后台提示词数组含固定时间字段索引，新增规则改为末尾追加，避免时间感知关闭时串字段。VM 测试补载新函数；旧网页英文放行预期按新授权改为拦截，不改变私人拒绝优先级。',
        '双版本线上、约会、共同生活、电话开关两态 16 组合，后台无气泡/无通知/无动作与确认、带翻译语音、授权笔记核心聊天门禁、原文多场景及设置持久化通过。全量最终结果见 .qa/english-policy-tests.xml。空回复具体用户现场仍未确定，本轮没有修改该问题。',
    ],
    'AI开发项目_Bug修改规范.docx': [
        '正文语言过滤必须与原文开关独立，但不能全局套在所有模型请求上，避免误拦总结、模型连通测试、图像提示词或 JSON 字段。判断时排除心情和功能元数据，合法译文必须能参与正文判定。拦截须早于动作及通知；手动释放和后台旧队列也须核对。',
        '网页与私人入口可能同名不同路径，必须以实际请求与可见结果验证覆盖。只加强标签提示词不能宣称已保证模型永不格式出错；不能顺手修复用户明确暂缓的空回复、替换心情解析器或扩大中文出戏策略。',
    ],
}

if __name__ == '__main__':
    for name, rows in ROWS.items():
        target = ROOT / 'docs' / 'maintenance' / name
        doc = Document(target)
        old = [p.text for p in doc.paragraphs]
        assert TITLE not in old, 'Do not append this entry twice'
        doc.add_page_break()
        doc.add_paragraph(TITLE, style='Heading 1')
        for row in rows:
            doc.add_paragraph(row)
        doc.save(target)
        assert [p.text for p in Document(target).paragraphs][:len(old)] == old
        print(name + ': appended, history preserved')
