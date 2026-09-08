"""Append the unshipped reply fixes without changing existing maintenance history."""
from pathlib import Path
from docx import Document
from docx.shared import Pt

ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / 'docs' / 'maintenance'
MARKER = '2026年9月8日 消息与小事簿修复 本地未发布'
COMMON = [
    '状态与范围：两边共有，网页 app.js、cohab-theater.js 与私人 PhoneWeb.bundle 对应逻辑分别修改，保留私人原有性能与恢复实现；后台 phone-role-push 仅改本地源码。用户要求先完成修改，等待明确授权再打包。本次没有提交、推送、部署服务器或生成安装包，不提升或复用发布版本号；正式发布必须重新分配唯一网页版本并核对私人资源包含关系。',
    '已验证与限制：完全隔离外网的浏览器夹具覆盖网页和私人网页入口的开关关闭、开启、再次关闭，微信、电话字幕、单次约会、共同生活、多人配角、合法内心/发推/来电标签及设置持久化。新增 Node 运行时测试覆盖 X 异步角色和账号归属、小事簿证据与落地时序、后台原文及待回复/锁定边界、查手机后回复。没有调用真实模型、真实 Supabase、真实支付或真实设备，不称 iPhone 音频播放及后台定时行为已验收。全量测试最终结果见同目录消息与小事簿现场记录的收尾章节。',
]
RECORDS = {
    'AI开发项目_项目说明文档.docx': [
        '新增设置入口为设置中的“角色、时间与消息”及偏好页“模型原文输出（实验）”，默认关闭。开启只绕过角色正文的格式、重复、旁白和措辞后处理；合法功能标签照常执行，未知正文不因格式被抛弃。关闭恢复原处理，不覆盖用户模型路线、密钥、语音配置或历史消息。接口没有返回正文仍不能制造回复。',
        '小事簿只记录值得记住的情绪变化、已明确情绪原因及已经答应的约定。停止用户发一句就按关键词搬运截断的旧自动记录，改为同次角色回复附带证据与角色视角总结。情绪证据来自用户；角色承诺必须在本轮实际进入对话记录后才可记入。愿望、反问、普通健康睡眠工作流水账不单独记录。新记录标明角色和账号归属；旧记录保留，不宣称历史自动文本为用户日记。',
        '主动消息补入最近用户消息之后角色已经说过的话，提示电量未变或经过时长本身不是新进展，不推断用户没回复就是躲避；不新增全局语义硬拦截。X 自动生成明确采用发布角色的独立模型路线，发布前再次核对角色对象、角色 ID 和账号归属，失效则不发布，不回落成另一角色。',
    ],
    'AI开发项目_Bug记录模板.docx': [
        '用户报告：偶发主动消息换词重复、电量与未回复话题再问；X 原应 A 发却似由 B 发，用户只能确认可能是定时自动；小事簿照搬碎句并误称用户自写；模型返回经应用后处理后无可显示正文。主动消息和 X 均无当次请求日志，不能将风险路径说成已复现的唯一根因。',
        '可确认根因：旧 lifeNoteOnUserMsg 按关键词取用户整句，交给 lifeNoteAdd 改代词和截短，可直接解释碎句与照抄。改为有证据的角色总结，并在实际消息入列时核验角色承诺，不在请求返回时提前记入。原有记录与自动记录开关保持可回退。',
        '排查与修复：X 独立请求原先未固定角色模型路线，增加 independentRoleModel 与路由参数并锁定异步归属；主动重复只补生成上下文，不放大全局相似度阻断。原文开关使用本轮快照覆盖微信、线下、通话、后台正文和查手机完成回复；仍保留任务 ID 幂等、账号会话过期、真实读取权限、设备/门锁授权及付款校验。',
        '测试中的失败与处理：旧源码正则和孤立 mock 因新增分支失败，更新合同断言及依赖，同时增加真正运行时用例。浏览器夹具先误调用多人场景闭包内部函数，后改为真实 offAI 入口；预览模式每次启动重新造种子，不能测持久化，现保存后进入普通 URL 验证。实页回归发现私人旧入口仍调用 lineToMsgs 压缩空格，补原文分支，不用网页整文件覆盖私人实现。',
        '并行现场风险：另一任务修改私人 HomeKit A100、原生版本、资源与 staging；其中私人 app.js 和剧场文件一度回到原基线，已只补回本次差异并复测。未修改、撤销或提交其他任务的原生文件；资源清单与新包必须由获得用户打包授权后的发布过程重新核对。',
    ],
    'AI开发项目_Bug修改规范.docx': [
        '原文实验不是取消功能权限。正文重复与格式过滤可关闭，真实设备授权、支付校验、会话/账号归属、后台消息或任务 ID 去重不得关闭。配置快照只传入需要展示给用户的角色请求，不能全局放开 JSON 动作决策、外卖、图片规划等辅助请求。后台功能须显式同步配置并另行部署，前端本地修改不等于已影响线上后台。',
        '涉及自主角色动作，生成入口必须固定角色人设和其独立模型路线；等待模型和媒体期间切账号、删除角色、替换对象时不得发布到新角色。不用当前页面选中人或错误回落弥补迟到结果。',
        '隐藏总结必须有本轮证据、正确说话者和作者归属。不能把用户愿望记成角色承诺；角色未实际说出的候选内容不得作为已承诺事实。只在消息已进入对应对话后收录，保留旧记录，不以未经核实的旧来源声称用户亲写。',
        '修改共享回复管线前列出网页与私人消费者，分别检查旧实现差异，不能把网页整文件覆盖私人性能或恢复逻辑。回归同时跑默认关闭与开启，并覆盖合法标签、未知正文、长重复文本、后台未完成回复、异步切号以及持久化。并行打包或 staging 可能覆盖本次逻辑，收尾必须复查实际文件，不能凭一次测试声明当前产物仍一致。',
    ],
}

def main():
    for filename, paragraphs in RECORDS.items():
        path = DOCS / filename
        doc = Document(path)
        old_text = [p.text for p in doc.paragraphs]
        if any(MARKER in p for p in old_text):
            raise RuntimeError(f'{filename}: already appended; do not overwrite history')
        old_tables = len(doc.tables)
        doc.add_page_break()
        heading = doc.add_paragraph(MARKER, style='Heading 1')
        for run in heading.runs:
            run.font.size = Pt(15)
        for text in [COMMON[0], *paragraphs, COMMON[1]]:
            doc.add_paragraph(text)
        doc.save(path)
        check = Document(path)
        assert [p.text for p in check.paragraphs][:len(old_text)] == old_text
        assert len(check.tables) == old_tables
        assert sum(MARKER in p.text for p in check.paragraphs) == 1
        print(f'{filename}: appended; {len(old_text)} original paragraphs and {old_tables} tables preserved')

if __name__ == '__main__':
    main()
