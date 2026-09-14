from io import BytesIO
from pathlib import Path
import subprocess
import zipfile

from docx import Document

ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs/maintenance"
TITLE = "v1241 和 v1242 大存档备份流畅修复"
FILES = [
    "AI开发项目_项目说明文档.docx",
    "AI开发项目_Bug记录模板.docx",
    "AI开发项目_Bug修改规范.docx",
]

for name in FILES:
    path = DOCS / name
    rel = path.relative_to(ROOT).as_posix()
    before_bytes = subprocess.check_output(["git", "show", f"HEAD:{rel}"], cwd=ROOT)
    before = [p.text for p in Document(BytesIO(before_bytes)).paragraphs]
    current_doc = Document(path)
    current = [p.text for p in current_doc.paragraphs]
    assert current[: len(before)] == before, f"{name}: historical paragraphs changed"
    assert current.count(TITLE) == 1, f"{name}: release heading missing or duplicated"
    assert len(current) == len(before) + 5, f"{name}: unexpected paragraph count"
    assert len(current_doc.sections) >= 1
    with zipfile.ZipFile(path) as archive:
        assert archive.testzip() is None
        assert "word/document.xml" in archive.namelist()
    print(f"verified {name}: {len(before)} historical paragraphs preserved, 5 appended")
