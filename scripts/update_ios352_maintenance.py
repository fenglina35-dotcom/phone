"""Append native-only Screen Time evidence without replacing prior records."""
from pathlib import Path
from docx import Document
from lxml import etree

root = Path(__file__).resolve().parents[1] / 'docs' / 'maintenance'
title = '私人 iOS352 屏幕同步修复候选 2026年9月10日'
updates = {
    'AI开发项目_项目说明文档.docx': '仅私人原生升级至1.0.352（352），内置网页及公开网页保持v1229、桥37。先保存本轮报告请求再刷新真实同步页的可见报告；角色本地检查与上传共用一轮8秒有界读取，避免请求编号互相覆盖。增加复制屏幕同步诊断，不改授权、已选App、共享组、锁定、健康与智能家居。日常事件簿、语音及此前私人功能完整保留。',
    'AI开发项目_Bug记录模板.docx': '用户报告有总时长显示但报告未回传、逐App不可用。对比9月8日9点前afe3b4cb，报告扩展实现与签名共享组没有改变，未发现公开版混入。确认旧可见视图先刷新、后经健康等待才创建请求，且不接收请求刷新通知；另有多个入口可覆盖共享请求。已修正两点，但这是源码确认的潜在失配原因，不是已证实的唯一真机根因。4项源码契约回归在旧实现3失败、新实现全通过；全仓1782项与网页及私人实际浏览器核心聊天通过。首次全量另遇零点附近拒接测试使用Date.now减一分钟跨日失败，未修改聊天代码，跨过零点后重跑通过。Swift编译、签名与真机同步尚未验证。',
    'AI开发项目_Bug修改规范.docx': '可见报告数字不代表主App已收到当前快照，更不代表云端同步成功。写请求必须先于触发刷新，多入口读取须合并，严格保留请求编号与时间校验，不用旧值或零值伪装成功。纠正历史资料中过强的全球回传假设：Apple DTS明确报告扩展沙箱限制导出；既有App Group兼容行为不能保证所有系统环境支持。缺少真机证据时标记修复候选，检查系统、Team、各Target签名与匿名诊断，不盲目重置权限或清空数据。',
}
for name, body in updates.items():
    path = root / name
    doc = Document(path)
    assert title not in [p.text for p in doc.paragraphs]
    old = [etree.tostring(x) for x in doc.element.body if not x.tag.endswith('}sectPr')]
    doc.add_page_break()
    doc.add_heading(title, 1)
    doc.add_paragraph(body)
    doc.add_paragraph('完整证据见私人_iOS352_屏幕同步修复核验.md。Apple DTS：https://developer.apple.com/forums/thread/817516 。打包与推送状态以实际发布结果为准；源码包不是可直接安装IPA。维护历史正文逐节点保留，捆绑LibreOffice缺失，文档版式未完成渲染验收。')
    doc.save(path)
    check = Document(path)
    assert [etree.tostring(x) for x in check.element.body if not x.tag.endswith('}sectPr')][:len(old)] == old
    print(name + ': appended; previous body preserved')
