"""Append scoped refusal-delivery findings; preserve every historical paragraph."""
from pathlib import Path
from docx import Document
ROOT=Path(__file__).resolve().parents[1]
TITLE='2026年9月8日 私人版后台模型拒绝说明隔离'
RECORDS={
 'AI开发项目_项目说明文档.docx':[
  '用户明确仅私人版：后台明确自称 AI 并解释安全政策的拒绝说明不作为角色对白，即使开启原样输出也不例外。北京时间气泡与快捷指令仍两边共有，未因此改变归属。',
  '私人客户端同步 privateBackgroundModelErrorGuard=true；共享后台只对严格启用标记的 profile 增加拒绝错误分支。私人收件增加旧队列防护和角色级错误记录。正常角色表达不同意、不愿意、换话题不受该规则影响，网页、前台聊天、电话和智能家居路径保持原规则。',
  '代码仅本地修改与隔离测试，未打包、推送、部署或真机验证。后台锁屏通知需部署函数并同步私人 profile，不能仅凭私人源码已改就宣称已生效。快捷指令后台是尚未完成的独立开发工作。'],
 'AI开发项目_Bug记录模板.docx':[
  '截图后台拒绝说明漏入角色气泡：完整原文在旧 roleMessage 实际复现，原样开、关两分支均返回 message；旧客户端拒绝检查也放行。现有 roleModelOutputLeak 主要检测推理和提示词泄露，不检测本类 AI 拒绝，且原样输出提前放行正文。触发模型拒绝的具体请求未取得，不能归因于用户新聊了敏感内容。',
  '私人标记下在后台原样放行及消息/动作投递之前识别明确拒绝文字、API refusal 字段或 content_filter/refusal 原因，作为 private-model-refusal 失败返回；同任务不自动重试，不要求模型绕过安全拒绝。私人收件另行防护旧队列，保存错误码/时刻/行 ID，先落盘再确认，不清理既有聊天。',
  '新增回归修复前 7 项中 6 项失败，修复后 7 项通过，包含实际漏投递重现而非仅缺函数。真实本地浏览器覆盖原文开关、动作不执行、正常对白、重复收件、错误持久化及 busy 释放；首次夹具漏掉私人容器能力标记而返回 false，补充夹具后通过，没有改变容器权限。',
  '核心聊天浏览器门禁及原样输出多场景回归通过，全量 Node 1634/1634。并行北京时间任务使旧源码 timestamp 断言失效，已核验关闭回退并补 UTC+8 数据和浏览器可见测试，非后台拒绝补丁引入的业务错误。测试未调用真实模型或真实设备；详见 私人版后台模型拒绝说明隔离_2026-09-08.md。'],
 'AI开发项目_Bug修改规范.docx':[
  '私人专属策略不得通过共享服务器默认分支影响网页：由私人入口明确同步标记，后端严格按 true 执行；同时测试无标记旧 profile 与网页原行为。新后台防护必须检查生成、动作、通知、收件、确认和持久化，不可只改气泡渲染。',
  '区分技术调用失败与角色本人的拒绝。明确的模型安全拒绝保留系统诊断，不制造角色台词，不以无限重试、供应商轮换或提示模型改写作为绕过手段。缺少触发请求时，分别记录已复现的漏投递原因和未确认的模型拒绝原因。']}
for name,rows in RECORDS.items():
 p=ROOT/'docs/maintenance'/name;d=Document(p);old=[x.text for x in d.paragraphs]
 assert TITLE not in old
 d.add_page_break();d.add_paragraph(TITLE,style='Heading 1')
 for row in rows:d.add_paragraph(row)
 d.save(p);assert [x.text for x in Document(p).paragraphs][:len(old)]==old
 print(name+': appended, history preserved')
