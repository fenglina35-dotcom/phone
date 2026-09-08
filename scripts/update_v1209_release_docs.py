"""Append the authorized combined release, preserving existing document history."""
from pathlib import Path
from docx import Document
ROOT=Path(__file__).resolve().parents[1]
TITLE='2026年9月8日 v1209 分离网页与私人功能的汇总发布'
RECORDS={
 'AI开发项目_项目说明文档.docx':[
  '用户已授权将消息修复、智能家居与卡顿任务的当前成果一起打包推送，取代上一阶段等待打包授权的状态。网页 v1209；私人网页 v1209；私人 iOS 1.0.333（333）；原生桥36。两个旧 v1128 私人工作树保留为历史现场，不做旧整树合并或清理。',
  '消息续聊、X 角色归属、小事簿情绪和约定、原文输出实验属于两边共有。智能家居新门锁、独立设备选择页、私人后台任务单通道与任务耗时诊断只进入私人源码及私人网页资源。网页版不包含新门锁脚本和调度器，原网页 smart-home.js 不变。',
  '私人包保留原小灯入口，门锁独立页面可自定义显示名称。关锁和解锁均须已知锁态与执行后回读；每次解锁必须 Face ID。强退恢复只取得观察时间时，不编造事件时间。卡顿修复只串行化自动任务，手动聊天、电话和闹钟入口不经过自动队列。',
 ],
 'AI开发项目_Bug记录模板.docx':[
  '发布核验发现私人门锁运行目录有 private-smart-lock.js/css，但被旧 PhoneWeb.bundle 通配忽略规则排除；这会导致工作区测试通过而从 Git 打包缺文件。已显式跟踪两份资源，并核对 Resources/Web 源和包内运行文件一致。私人 web-hotfix.js 同样明确入库，但不改变私人已有入口加载策略。',
  '版本提升时旧测试同时核验当前说明和历史 iOS327 说明，机械更新错误触及历史路径及预期。已保留历史 v1206 文档路径和预期，只更新当前 v1209/iOS333 说明，没有修改历史事实。',
  '整包预检按最终跟踪文件清单检查脚本引用、网页与私人共享资源、像素衣柜全部子资源、门锁文件、12 个 Target 版本、双入口一致及网页排除私人标记。包必须从干净提交生成，再解包逐项比较。真实卡顿和当次偶发消息事件仍未复现；没有真实门锁动作或真机安装证据。',
 ],
 'AI开发项目_Bug修改规范.docx':[
  '跨任务发布以当前有效成果为单位，不以“所有工作区”作为把旧整树覆盖最新版本的理由。私人源码可与网页源码一同入仓，但网页入口和静态网页包必须不包含私人功能。仓库公开性与运行功能隔离是两件事，不将 native 目录误称保密。',
  '打包前检查 ignored-but-required 文件。工作区文件存在和本地测试通过，不能代替 git 跟踪清单及最终 ZIP 的包含证明。所有用户可见共享脚本、样式、嵌套素材和双入口必须反向核对；保留私人核心的性能、恢复及原生差异，禁止用旧 staging 脚本从网页整包覆盖私人核心。',
  '新网页和私人产物都使用唯一新版本、新包名。Windows 产物仅为 Mac 待编译源码候选，尚无 Mac 编译签名和真实 iPhone 验收，不可称可运行私人交付。推送只自动尝试一次，线上 Pages 与单独的后台函数部署分别核实，不混淆提交、包生成与上线状态。',
 ]}
for name,rows in RECORDS.items():
 p=ROOT/'docs/maintenance'/name;d=Document(p);old=[x.text for x in d.paragraphs]
 assert TITLE not in old
 d.add_page_break();d.add_paragraph(TITLE,style='Heading 1')
 for row in rows:d.add_paragraph(row)
 d.save(p);check=Document(p);assert [x.text for x in check.paragraphs][:len(old)]==old
 print(name+': appended; original history preserved')
