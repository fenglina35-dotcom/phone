"""Append release evidence while preserving every prior maintenance paragraph."""
from pathlib import Path
from docx import Document
from lxml import etree
root=Path(__file__).resolve().parents[1]/'docs'/'maintenance'
title='v1229 语音与日常事件簿汇总 2026年9月9日'
updates={
 'AI开发项目_项目说明文档.docx':'网页与私人内置v1229、私人iOS351、桥37，汇总v1228未发布语音修复。日常事件簿入口在查他手机，使用记仇本相同线条图标；默认关闭，100条可调50至300条。私人包保留全部共享功能、智能灯门锁空调及卡顿修复，公开North现场不打包，私人智能家居不进入网页。',
 'AI开发项目_Bug记录模板.docx':'旧v1227实际ZIP包含事件簿脚本、HTML加载及入口，未证实文件漏包；SPYICON缺events映射导致文字图标，已最小修复两端。新增真实页面断言先在旧实现失败，新实现相同SVG、空手机入口、实际点击、开关和上限保存通过。此前语音44场景及核心聊天双端回归通过，全仓1778项通过。测试夹具曾漏情侣绑定及误用预览重载，已仅修夹具，不改变产品权限和存档。',
 'AI开发项目_Bug修改规范.docx':'判断漏包应核验用户对应ZIP的脚本、入口引用、UI路由和实际运行版本，不用源码存在代替包内证据。图标复用既有SVG映射；新包须从提交HEAD制作，逐文件哈希验证并用解压后的入口进行实际浏览器检查。安装说明需明确当前版含哪些功能，将旧不适用说明标为历史。',
}
for name,body in updates.items():
 p=root/name;d=Document(p)
 assert title not in [x.text for x in d.paragraphs],name+' already appended'
 old=[etree.tostring(x) for x in d.element.body if not x.tag.endswith('}sectPr')]
 d.add_page_break();d.add_heading(title,1);d.add_paragraph(body)
 d.add_paragraph('详见v1229_语音与事件簿发布核验.md。实际打包推送结果另记，未完成前不能称已上线。模型、音频及原生传输为隔离测试，Mac签名编译、真实iPhone和真实声音未验证；维护文档已保留历史正文，缺少捆绑LibreOffice，版式渲染未验收。')
 d.save(p);r=Document(p)
 assert [etree.tostring(x) for x in r.element.body if not x.tag.endswith('}sectPr')][:len(old)]==old
 print(name+': appended, prior body preserved')
