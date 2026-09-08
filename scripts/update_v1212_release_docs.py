"""Append only; preserve all existing maintenance history."""
from pathlib import Path
from docx import Document
ROOT=Path(__file__).resolve().parents[1]
TITLE='2026年9月8日 v1212 请求故障诊断与简洁时间'
ROWS={
 'AI开发项目_项目说明文档.docx':[
  '两边共有：本机请求故障诊断，角色设置的模型与路线诊断内可查看和复制。开始即记记录，区分空 HTTP 正文、非 JSON、流式数据、缺少字段、空正文字段和请求错误，并将取得正文与聊天实际处理结果分开。记录本地限80行，响应片段仅页面内存，不进入复制报告，已配置密钥脱敏。诊断不额外调用模型、不改变原重试和原输出策略。',
  '北京时间仅显示时分秒，不显示年月日或北京时间文字，不撑宽气泡；消息原始时间与关闭回退不变。此前纯英文无译文拦截、心情格式提示两态覆盖一起发布。私人包另含9defbd8b门锁内部标签隐藏修复，Face ID和真实回读保留，网页不包含。v1212、私人iOS1.0.338（338）、桥36。未完成快捷指令草稿不提交、不入包、不部署。'],
 'AI开发项目_Bug记录模板.docx':[
  '受影响用户的真实空回复根因仍未确认，不能把计费摘要、输入输出计数或流式标记当成原始响应。已确认旧解析失败吞为null且只读取标准正文字段，导致空回复不可定位。新增观察保留HTTP、类型、请求ID、结束原因、规模、最后阶段及轮次结果。旧代码故障注入与真实浏览器空回复无记录断言失败，新实现通过；不宣称已修复真实上游。',
  '开发中无锚点补丁新增误落文件末尾、私人请求函数形态差异被语法和加载顺序测试捕获，已改为定点修改。身份升级时缓存重载键与脚本严格启动标记漏更新由版本回归捕获并修正。历史文档保持，不将旧安装记录改成新版本。',
  '双入口故障注入、已有授权及笔记核心聊天、原文开关/电话/约会/共同生活/后台隔离测试完成，短气泡开启前后同为37像素。最终全量测试、产物哈希与部署以v1212发布结果为准。Windows模拟HTTP不代替真实模型、签名或iPhone验收。'],
 'AI开发项目_Bug修改规范.docx':[
  '请求日志必须在发送前建立，失败和空响应保留原因；只记录成功结果不足以诊断。传输成功、正文提取成功和实际角色送达是不同阶段，不得互相替代。副作用观察不可增加调用或改变错误回退，存储配额失败不能中断核心聊天。复制诊断采用最小元数据，不自动包含聊天正文、响应片段和密钥。',
  '时间标签不能参与气泡正文的内在宽度计算。共享代码新增要核对私人真实入口，机械补丁必须带准确上下文；新增JS须显式入库并检查双入口和Service Worker引用。新版本同时核对所有资源查询、启动身份和恢复入口，再从干净提交打包。']}
for name,rows in ROWS.items():
 p=ROOT/'docs/maintenance'/name;d=Document(p);old=[x.text for x in d.paragraphs];assert TITLE not in old
 d.add_page_break();d.add_paragraph(TITLE,style='Heading 1')
 for text in rows:d.add_paragraph(text)
 d.save(p);assert [x.text for x in Document(p).paragraphs][:len(old)]==old
 print(name+': appended, old paragraphs retained')
