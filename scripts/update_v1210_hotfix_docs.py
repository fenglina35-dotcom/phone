"""Append the core chat regression and release gate without rewriting history."""
from pathlib import Path
from docx import Document
ROOT=Path(__file__).resolve().parents[1]
TITLE='2026年9月8日 v1210 核心聊天中断热修复'
RECORDS={
 'AI开发项目_项目说明文档.docx':[
  '用户要求立即修复网页和私人版的 v1209 聊天中断。两边共有：myActivity 将错误的未定义 c 改为已有角色变量 self。新网页与私人内置网页 v1210，私人 iOS 1.0.334（334），原生桥36不变；保留私人智能家居和后台卡顿修复，不进入网页运行资源。',
  '不清理聊天、小事簿、用户模型或语音配置，不撤销查手机授权，不改变过滤、去重、持久化或原生流程。私人产物仍为 Mac 待编译源码候选，不等于真机已安装；线上与包状态以本轮实际收尾核验为准。'],
 'AI开发项目_Bug记录模板.docx':[
  "P0：用户更新 v1209 后发消息显示模型未回复 Can't find variable: c。网页与私人入口均复现：getSpy(role).granted 为真且未通过 memorySince 跳过时，buildSystem 调用 myActivity，在模型请求前抛 ReferenceError。myActivity 已定义 self，却在小事簿筛选和作者说明误用 c。该错误由 v1209 小事簿改动引入，不能归因于模型拒答。没有全体用户遥测，不把所有用户的数量作为已验证事实。",
  '定点修复两份 app.js 的两处引用；不添加全局 c 或用异常吞掉上下文。此前回归默认未开启 role.spy.granted，只测了未授权分支，源码/语法检查无法发现运行时自由变量。新增浏览器脚本先在旧代码复现，再验证完整上下文、模型请求、回复可见送达及忙碌释放。',
  '新增 tests/chat-authorized-phone.test.mjs 与 scripts/check_chat_authorized_phone.cjs 覆盖授权开关、原文开关、空与已有小事簿、角色与账号隔离、时间增量、刷新、自动查看、真实输入框与让TA回、持久化重载。孤立VM缺失角色用例补 msgs/lastMsg 依赖；自动查看夹具先复用了上轮相同回复被正常去重，改模拟回复后通过，没有因此改业务去重。',
  '两边浏览器16组聊天组合、自动查看和持久化通过；模型HTTP使用隔离测试响应，未调用用户真实模型、真实设备或执行Mac签名。完整测试及最终部署记录见同目录 v1210_聊天中断热修复.md。'],
 'AI开发项目_Bug修改规范.docx':[
  '用户于此次严重事故后明确要求保存长期规则：发布必须保护核心功能。每次修改先界定共享消费者；核心聊天必须跑真实入口并覆盖已有权限和已有存档、角色与账号隔离、默认及实验模式、可见送达、忙碌释放与状态恢复。测试必须先能捕获原故障，不能只凭语法通过、源码正则、默认空存档或测试总数宣布可发布。',
  '变量重命名和新上下文引用必须在实际函数作用域验证。共享函数的普通聊天、电话上下文、手动刷新及自动查看都要列入回归，网页和私人入口分别检查。最终包与线上资源另行反向验证；无法验证的真机路径必须披露，不能以绝对无Bug承诺替代证据。']}
for name,rows in RECORDS.items():
 p=ROOT/'docs/maintenance'/name;d=Document(p);old=[x.text for x in d.paragraphs]
 assert TITLE not in old
 d.add_page_break();d.add_paragraph(TITLE,style='Heading 1')
 for row in rows:d.add_paragraph(row)
 d.save(p);assert [x.text for x in Document(p).paragraphs][:len(old)]==old
 print(name+': appended, history preserved')
