"""Append voice regression evidence without rewriting maintained history."""
from pathlib import Path
from docx import Document
from lxml import etree

ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / 'docs' / 'maintenance'
TITLE = 'v1228 语音格式回归核验 2026年9月9日'
DETAIL = '详见 docs/maintenance/v1228_语音格式回归核验_未发布.md。本轮仅本地候选，未提交、打包或推送；真实模型与语音服务、Mac编译签名及真实iPhone播放尚未验证。'
UPDATES = {
 'AI开发项目_项目说明文档.docx': [
  '网页及私人内置页共同候选v1228。明确请求和角色自主语音统一格式处理，原样输出开关两态均适用。合规回复不额外调用模型；缺语音格式或外语译文时最多纠正一次，可能产生一次模型调用费用。自动语音频率0仍不主动发语音，明确要求语音可以覆盖该频率。',
  '保持9月5日的语音接口、音色、语言和服务配置；不修改公开North、空调或智能家居。私人原生仍iOS350桥37，打包前另行更新包版本与说明。', DETAIL],
 'AI开发项目_Bug记录模板.docx': [
  '症状：先生你发句语音未被识别；原样输出模式下主动外语语音缺少翻译时被降级为纯英文文字。真实两端浏览器基线已复现。缺少中文译文的语音示例和把隐藏语气当作中文的判断会进一步掩盖该问题。',
  '历史对比：9月5日最后提交90accb87的标签解析、语言配置和ttsArr与本轮起点相同，旧版也漏识别该无逗号请求并存在缺译文降级。v1226仅改频率覆盖及兜底未覆盖原样输出完整流程；本轮使用真实aiReply消息落库回归，不重复仅凭辅助函数或正则断言宣布修复。',
  '修复：常见称呼无需逗号；两态语音候选最终统一校验与一次纠正；隐藏语气不算译文；失败诊断不得降成英文正文。根网页和私人审计上下文不同，必须分别接入；开发中双入口回归已捕获并修正私人变量引用差异，失败候选未发布。',
  '验证：新增9项真实函数测试，全仓1778项通过；实际聊天回归44场景通过，授权查手机后的普通聊天、角色账号隔离、存档重载通过。音频预热为测试替身，不冒充真实声音已播放。', DETAIL],
 'AI开发项目_Bug修改规范.docx': [
  '语音修复必须覆盖截图原话及无标点称呼、主动与明确两入口、原样输出两态、自动频率0、缺译文、纠正仍失败、普通聊天和否定语句。测试必须使用实际标签解析与消息落库，不用返回空对象的假解析器证明语音可用。',
  '旧版可用是现场线索而非整版正确的证明。逐函数核对历史版本，明确哪些函数未变、哪些后处理改变及设置变量差异。保留TTS base/key/model/voice，不伪造翻译和语音成功；修复请求有界并告知额外成本。',
  '跨网页/私人同步须逐一核对作用域和现有诊断包装层，不能假设同名变量都存在；先语法检查，再实际双入口聊天，失败即阻止发布。', DETAIL],
}
for filename, paragraphs in UPDATES.items():
    path = DOCS / filename
    doc = Document(path)
    if TITLE in [p.text for p in doc.paragraphs]:
        raise RuntimeError('already appended: ' + filename)
    old = [etree.tostring(el) for el in doc.element.body if not el.tag.endswith('}sectPr')]
    doc.add_page_break()
    doc.add_heading(TITLE, 1)
    for paragraph in paragraphs:
        doc.add_paragraph(paragraph)
    doc.save(path)
    reread = Document(path)
    current = [etree.tostring(el) for el in reread.element.body if not el.tag.endswith('}sectPr')]
    assert current[:len(old)] == old, 'history changed: ' + filename
    print(filename + ': appended; prior body preserved')
